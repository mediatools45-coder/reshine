"""Builds the restoration step guide from docs/guide/steps.json + notes.json + img/:
docs/guide/guide.md (one section per object, used to build the Google Doc tabs) and
docs/guide/Restoration_Step_Guide.docx (one heading + page per object)."""
import json, glob, os
from docx import Document
from docx.shared import Cm, Pt
from docx.enum.text import WD_BREAK
G = 'docs/guide'
d = json.load(open(f'{G}/steps.json'))['levels']; n = json.load(open(f'{G}/notes.json'))
def tool_of(s):
    if s['tool'] in ('hand', 'none'): return None, n['handLabels'][s['toolLabel']]
    t = n['tools'][s['tool']]
    if s['toolLabel'] and s['toolLabel'].lower() not in t.lower(): t = f"{s['toolLabel']} ({t.split(':')[0].lower()}):{t.split(':', 1)[1]}"
    return f"img/tools/{s['tool']}.jpg", t
md = ['# Restoration Workshop: step guide', '', 'One section per object (one tab each in the Google Doc). Each step has a picture of the object, the tool and one line on what the tool does.', '']
doc = Document(); doc.styles['Normal'].font.size = Pt(10.5)
doc.add_heading('Restoration Workshop: step guide', 0)
doc.add_paragraph('One section per object. Each step has a picture of the object, the tool and one line on what the tool does. Use the Navigation pane (View > Navigation Pane) to jump between objects.')
for k, v in d.items():
    md += [f"## {v['name']}", '']
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE); doc.add_heading(v['name'], 1)
    shots = sorted(glob.glob(f'{G}/img/{k}/*.jpg'))
    for i, s in enumerate(v['steps']):
        img, tnote = tool_of(s); snote = n['steps'][k][s['id']]; shot = os.path.relpath(shots[i], G)
        md += [f"### Step {i + 1}: {s['name']}", '', snote, '', f"![{s['name']}]({shot})"] + ([f"![tool]({img})"] if img else []) + ['', f"**{tnote}**", '']
        h = doc.add_heading(f"Step {i + 1}: {s['name']}", 2); h.paragraph_format.keep_with_next = True
        doc.add_paragraph(snote).paragraph_format.keep_with_next = True
        t = doc.add_table(rows=1, cols=2); c0, c1 = t.rows[0].cells
        c0.paragraphs[0].add_run().add_picture(shots[i], width=Cm(7.5))
        if img: c1.paragraphs[0].add_run().add_picture(f'{G}/{img}', width=Cm(5))
        p = c1.add_paragraph(); r = p.add_run(tnote); r.bold = True
open(f'{G}/guide.md', 'w').write('\n'.join(md))
doc.save(f'{G}/Restoration_Step_Guide.docx'); print('ok')
